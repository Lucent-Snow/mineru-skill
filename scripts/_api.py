#!/usr/bin/env python3
"""
MinerU API 核心封装
提供文件验证、上传、轮询、下载等功能，供 process_document / process_batch 调用。
"""
import json
import asyncio
import random
import time
import zipfile
import shutil
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from niquests import AsyncSession

from _config import (
    TOKENS_FILE, API_BASE_URL, MAX_FILE_SIZE, MAX_PAGES,
    SUPPORTED_FORMATS, MODEL_AUTO_SELECT,
)


# ============================================================
# Token 管理
# ============================================================

def load_tokens() -> Dict:
    """加载 all_tokens.json，返回 {email: {token, name, ...}} 字典"""
    try:
        with open(TOKENS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return {}


def get_random_token(tokens: Dict) -> str:
    """随机选择一个 Token（负载均衡）"""
    email = random.choice(list(tokens.keys()))
    return tokens[email]["token"]


def check_tokens_valid(tokens: Dict) -> Tuple[bool, List[str]]:
    """检查所有 Token 是否有效，返回 (全部有效, 过期邮箱列表)"""
    from datetime import datetime, timezone
    now = datetime.now(timezone.utc)
    expired = []
    for email, info in tokens.items():
        exp = datetime.fromisoformat(info["expired_at"].replace("Z", "+00:00"))
        if (exp - now).days <= 0:
            expired.append(email)
    return len(expired) == 0, expired


# ============================================================
# 文件验证
# ============================================================

def validate_file(file_path: str) -> Tuple[bool, str, Dict]:
    """
    验证本地文件：存在性、大小、格式、页数。
    返回 (有效, 错误信息, 文件信息字典)
    """
    path = Path(file_path)

    if not path.exists():
        return False, "文件不存在", {}

    size = path.stat().st_size
    if size > MAX_FILE_SIZE:
        return False, f"文件超过200MB限制 ({size / 1024 / 1024:.1f}MB)", {}
    if size == 0:
        return False, "文件为空", {}

    suffix = path.suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_FORMATS:
        return False, f"不支持的格式: {suffix}", {}

    pages = _get_page_count(str(path), suffix)

    return True, "", {
        "path": str(path),
        "name": path.name,
        "size": size,
        "format": suffix,
        "is_url": False,
        "pages": pages,
        "needs_split": pages > MAX_PAGES if pages else False,
    }


async def validate_url(session: AsyncSession, url: str) -> Tuple[bool, str, Dict]:
    """
    验证 URL 文件：可访问性、大小、格式（含 magic bytes 识别）。
    返回 (有效, 错误信息, 文件信息字典)
    """
    try:
        response = await session.head(url, timeout=10, allow_redirects=True)
        if response.status_code != 200:
            return False, f"URL无法访问: {response.status_code}", {}

        size = int(response.headers.get("content-length", 0))
        if size > MAX_FILE_SIZE:
            return False, f"文件超过200MB限制 ({size / 1024 / 1024:.1f}MB)", {}

        content_type = response.headers.get("content-type", "")
        fmt = _guess_format_from_url(url, content_type)

        # Fallback: magic bytes
        if not fmt:
            resp = await session.get(url, timeout=10, headers={"Range": "bytes=0-16"})
            head = resp.content[:16]
            fmt = _detect_format_by_magic(head, url)

        if not fmt:
            return False, "无法识别文件格式", {}

        return True, "", {
            "path": url,
            "name": Path(url).name or "document",
            "size": size,
            "format": fmt,
            "is_url": True,
            "pages": None,
            "needs_split": False,
        }
    except Exception as e:
        return False, f"URL验证失败: {e}", {}


# ============================================================
# 核心 API 操作
# ============================================================

async def upload_file(
    session: AsyncSession, token: str, file_path: str, **options
) -> Optional[str]:
    """
    上传本地文件到 MinerU。
    返回 batch_id，失败返回 None。
    """
    headers = {
        "authorization": f"Bearer {token}",
        "content-type": "application/json",
    }
    file_name = Path(file_path).name

    # 获取上传链接
    data = {"files": [{"name": file_name}], **options}
    response = await session.post(
        f"{API_BASE_URL}/file-urls/batch",
        headers=headers,
        json=data,
        timeout=30,
    )
    result = response.json()

    if result["code"] != 0:
        print(f"  获取上传链接失败: {result.get('msg')}")
        return None

    batch_id = result["data"]["batch_id"]
    upload_url = result["data"]["file_urls"][0]

    # 上传文件二进制
    with open(file_path, "rb") as f:
        file_data = f.read()

    upload_resp = await session.put(upload_url, data=file_data, timeout=300)
    if upload_resp.status_code == 200:
        return batch_id

    print(f"  文件上传失败: HTTP {upload_resp.status_code}")
    return None


async def wait_for_completion(
    session: AsyncSession, token: str, batch_id: str, max_wait: int = 600
) -> Optional[List[Dict]]:
    """
    轮询任务状态直到完成或超时。
    返回结果列表，失败/超时返回 None。
    """
    headers = {"authorization": f"Bearer {token}"}
    start = time.time()

    while time.time() - start < max_wait:
        resp = await session.get(
            f"{API_BASE_URL}/extract-results/batch/{batch_id}",
            headers=headers,
            timeout=30,
        )
        data = resp.json()

        if data["code"] != 0:
            await asyncio.sleep(5)
            continue

        results = data["data"]["extract_result"]
        all_done = True

        for r in results:
            state = r.get("state")
            if state == "failed":
                print(f"  任务失败: {r.get('err_msg')}")
                return None
            if state in ("pending", "running", "waiting-file", "converting"):
                all_done = False
                if state == "running":
                    progress = r.get("extract_progress", {})
                    extracted = progress.get("extracted_pages", 0)
                    total = progress.get("total_pages", 0)
                    if total > 0:
                        print(f"  处理进度: {extracted}/{total} 页", end="\r")

        if all_done:
            print()  # 换行
            return results

        await asyncio.sleep(5)

    print("  任务超时")
    return None


async def download_and_extract(
    session: AsyncSession, zip_url: str, output_dir: str
) -> Optional[str]:
    """
    下载结果 zip 并解压到指定目录。
    返回解压目录路径，失败返回 None。
    """
    try:
        response = await session.get(zip_url, timeout=300)
        if response.status_code != 200:
            print(f"  下载失败: HTTP {response.status_code}")
            return None

        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        zip_path = out / "result.zip"

        with open(zip_path, "wb") as f:
            f.write(response.content)

        with zipfile.ZipFile(zip_path, "r") as zf:
            zf.extractall(out)

        zip_path.unlink()
        return str(out)
    except Exception as e:
        print(f"  下载解压失败: {e}")
        return None


def find_markdown(directory: str) -> Optional[str]:
    """在目录中查找第一个 .md 文件"""
    for md in Path(directory).rglob("*.md"):
        return str(md)
    return None


def organize_output(extracted_dir: str, source_path: str, output_dir: str) -> Dict:
    """
    整理输出：将 Markdown 和图片复制到目标目录。
    返回 {markdown: 路径, images: 路径}
    """
    file_stem = Path(source_path).stem
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    md_file = out / f"{file_stem}.md"
    images_dir = out / f"{file_stem}_images"

    # 复制 Markdown
    source_md = find_markdown(extracted_dir)
    if source_md:
        shutil.copy(source_md, md_file)

    # 复制图片
    source_images = Path(extracted_dir) / "images"
    image_count = 0
    if source_images.exists():
        if images_dir.exists():
            shutil.rmtree(images_dir)
        shutil.copytree(source_images, images_dir)
        image_count = len(list(images_dir.glob("*")))

    return {
        "markdown": str(md_file) if md_file.exists() else None,
        "images": str(images_dir) if images_dir.exists() else None,
        "image_count": image_count,
    }


# ============================================================
# 内部辅助
# ============================================================

def _get_page_count(file_path: str, fmt: str) -> Optional[int]:
    """获取文件页数"""
    try:
        if fmt == "pdf":
            from PyPDF2 import PdfReader
            return len(PdfReader(file_path).pages)
        elif fmt in ("pptx", "ppt"):
            from pptx import Presentation
            return len(Presentation(file_path).slides)
        elif fmt in ("docx", "doc"):
            from docx import Document
            return len(Document(file_path).paragraphs) // 5
    except Exception:
        pass
    return None


def _guess_format_from_url(url: str, content_type: str) -> Optional[str]:
    """从 URL 路径或 Content-Type 推断格式"""
    url_lower = url.lower()
    for ext in SUPPORTED_FORMATS:
        if url_lower.endswith(f".{ext}"):
            return ext
    for ext, mime in SUPPORTED_FORMATS.items():
        if mime in content_type:
            return ext
    return None


def _detect_format_by_magic(head: bytes, url: str = "") -> Optional[str]:
    """通过文件魔数识别格式"""
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        url_lower = url.lower()
        if any(ext in url_lower for ext in ["pptx", "ppt"]):
            return "pptx"
        if any(ext in url_lower for ext in ["docx", "doc"]):
            return "docx"
        return "docx"  # 默认 ZIP-based 为 docx
    if head[:3] == b"\xff\xd8\xff":
        return "jpg"
    if head[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    lower = head.lower()
    if b"<html" in lower or b"<!doctype" in lower:
        return "html"
    return None


def is_url(path: str) -> bool:
    """判断路径是否为 URL"""
    return path.startswith(("http://", "https://"))
